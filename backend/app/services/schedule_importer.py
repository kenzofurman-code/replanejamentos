import xml.etree.ElementTree as ET
import openpyxl
import datetime
from typing import Dict, List, Any, Optional
import io
import unicodedata
import os
import tempfile
import jpype

_JVM_INITIALIZED = False

def _ensure_jvm():
    global _JVM_INITIALIZED
    if _JVM_INITIALIZED and jpype.isJVMStarted():
        return
    curr_dir = os.path.dirname(os.path.abspath(__file__))
    jvm_dll = os.path.abspath(os.path.join(curr_dir, "..", "..", "tools", "jre17", "bin", "server", "jvm.dll"))
    if not os.path.exists(jvm_dll):
        jvm_dll = jpype.getDefaultJVMPath()

    import mpxj
    mpxj_dir = os.path.join(os.path.dirname(mpxj.__file__), "lib")
    jars = [os.path.join(mpxj_dir, f) for f in os.listdir(mpxj_dir) if f.endswith(".jar")]
    if not jpype.isJVMStarted():
        jpype.startJVM(jvm_dll, classpath=jars)
    _JVM_INITIALIZED = True

def norm(text: str) -> str:
    if not text:
        return ""
    return unicodedata.normalize('NFKD', str(text)).encode('ASCII', 'ignore').decode('utf-8').lower().strip()

def try_float(val, default=0.0) -> float:
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    try:
        s = str(val).strip().replace(',', '.')
        return float(s) if s else default
    except (ValueError, TypeError):
        return default

def try_date(val) -> Optional[datetime.date]:
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    # Handle Java LocalDateTime / Date objects via JPype
    if hasattr(val, "getYear") and hasattr(val, "getMonthValue") and hasattr(val, "getDayOfMonth"):
        try:
            return datetime.date(int(val.getYear()), int(val.getMonthValue()), int(val.getDayOfMonth()))
        except Exception:
            pass
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    for fmt in ("%a %b %d %H:%M:%S %Z %Y", "%Y-%m-%dT%H:%M:%SZ"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def parse_ms_project_xml(xml_content: bytes) -> Dict[str, Any]:
    """
    Parses an MS Project XML file (.xml exported from MS Project).
    Extracts tasks, durations, start/end dates, predecessors, successors, lote_mae and lot.
    """
    root = ET.fromstring(xml_content)
    ns = {'mp': root.tag.split('}')[0].strip('{')} if '}' in root.tag else {}
    
    tasks_dict: Dict[str, Dict[str, Any]] = {}
    tasks_element = root.find('mp:Tasks', ns) if ns else root.find('Tasks')
    
    uid_to_name: Dict[str, str] = {}
    uid_to_id: Dict[str, str] = {}
    outline_to_name: Dict[str, str] = {}
    raw_tasks = []

    if tasks_element is not None:
        task_nodes = tasks_element.findall('mp:Task', ns) if ns else tasks_element.findall('Task')
        for t_node in task_nodes:
            name_el = t_node.find('mp:Name', ns) if ns else t_node.find('Name')
            if name_el is None or not name_el.text:
                continue
            name = name_el.text.strip()
            
            uid_el = t_node.find('mp:UID', ns) if ns else t_node.find('UID')
            id_el = t_node.find('mp:ID', ns) if ns else t_node.find('ID')
            outline_el = t_node.find('mp:OutlineNumber', ns) if ns else t_node.find('OutlineNumber')
            summary_el = t_node.find('mp:Summary', ns) if ns else t_node.find('Summary')
            
            uid = uid_el.text.strip() if uid_el is not None and uid_el.text else str(len(raw_tasks) + 1)
            t_id = id_el.text.strip() if id_el is not None and id_el.text else uid
            outline = outline_el.text.strip() if outline_el is not None and outline_el.text else ""
            wbs_el = t_node.find('mp:WBS', ns) if ns else t_node.find('WBS')
            wbs = wbs_el.text.strip() if wbs_el is not None and wbs_el.text else outline
            level_el = t_node.find('mp:OutlineLevel', ns) if ns else t_node.find('OutlineLevel')
            outline_level = int(level_el.text.strip()) if level_el is not None and level_el.text else (len(wbs.split('.')) if wbs else 1)
            is_summary = (summary_el is not None and summary_el.text in ('1', 'true', 'True'))

            # Ignore project root container (ID 0)
            if t_id == "0":
                continue

            uid_to_name[uid] = name
            uid_to_id[uid] = t_id
            if outline:
                outline_to_name[outline] = name

            # Duration (format PT...H...M... or days)
            dur_el = t_node.find('mp:Duration', ns) if ns else t_node.find('Duration')
            dur_days = 0.0
            if dur_el is not None and dur_el.text:
                dur_str = dur_el.text
                if 'H' in dur_str:
                    try:
                        h_part = dur_str.replace('PT', '').split('H')[0]
                        dur_days = float(h_part) / 8.0
                    except Exception:
                        dur_days = 1.0
                else:
                    dur_days = 1.0

            start_el = t_node.find('mp:Start', ns) if ns else t_node.find('Start')
            finish_el = t_node.find('mp:Finish', ns) if ns else t_node.find('Finish')
            
            start_date = try_date(start_el.text if start_el is not None else None)
            finish_date = try_date(finish_el.text if finish_el is not None else None)

            # Predecessors
            preds = []
            pred_nodes = t_node.findall('mp:PredecessorLink', ns) if ns else t_node.findall('PredecessorLink')
            for p in pred_nodes:
                p_uid_el = p.find('mp:PredecessorUID', ns) if ns else p.find('PredecessorUID')
                if p_uid_el is not None and p_uid_el.text:
                    p_uid = p_uid_el.text.strip()
                    type_el = p.find('mp:Type', ns) if ns else p.find('Type')
                    t_val = type_el.text.strip() if type_el is not None and type_el.text else "1"
                    type_map_xml = {'0': 'FF', '1': 'FS', '2': 'SF', '3': 'SS'}
                    rtype = type_map_xml.get(t_val, 'FS')
                    lag_el = p.find('mp:LinkLag', ns) if ns else p.find('LinkLag')
                    lag_val = 0
                    if lag_el is not None and lag_el.text:
                        try:
                            raw_lag = int(lag_el.text.strip())
                            lag_val = round(raw_lag / 4800) if abs(raw_lag) >= 4800 else raw_lag
                        except Exception:
                            pass
                    preds.append((p_uid, rtype, lag_val))

            raw_tasks.append({
                "uid": uid,
                "id": t_id,
                "outline": wbs,
                "wbs": wbs,
                "outline_level": outline_level,
                "name": name,
                "duration": dur_days,
                "start": start_date,
                "finish": finish_date,
                "pred_tuples": preds,
                "is_summary": is_summary
            })

    # Build successor relationships and daily values
    succ_map: Dict[str, List[str]] = {t["id"]: [] for t in raw_tasks}
    pred_map: Dict[str, List[str]] = {t["id"]: [] for t in raw_tasks}

    for t in raw_tasks:
        pred_formatted = []
        for p_uid, rtype, lag in t.get("pred_tuples", []):
            p_id = uid_to_id.get(p_uid, p_uid)
            lag_str = f'+{lag}d' if lag > 0 else (f'{lag}d' if lag < 0 else '')
            pred_formatted.append(f'{p_id}{rtype}{lag_str}')
            if p_id in succ_map:
                succ_map[p_id].append(t["id"])
        pred_map[t["id"]] = pred_formatted

    for t in raw_tasks:
        start_date = t["start"]
        finish_date = t["finish"]
        daily_vals: Dict[datetime.date, float] = {}

        if start_date and finish_date and finish_date >= start_date:
            total_days = max(1, (finish_date - start_date).days + 1)
            daily_pct = 100.0 / total_days
            curr = start_date
            while curr <= finish_date:
                if curr.weekday() < 5:
                    daily_vals[curr] = daily_pct
                else:
                    daily_vals[curr] = 0.0
                curr += datetime.timedelta(days=1)
            
            active_days = sum(1 for v in daily_vals.values() if v > 0)
            if active_days > 0:
                pct_per_day = 100.0 / active_days
                for d in daily_vals:
                    if daily_vals[d] > 0:
                        daily_vals[d] = pct_per_day

        # Derive lote_mae and lot from outline or task name
        parts = t["name"].split(" - ")
        if len(parts) >= 3:
            lote_mae = parts[1].strip()
            lot = " - ".join(parts[2:]).strip()
        elif len(parts) == 2:
            lote_mae = parts[0].strip()
            lot = parts[1].strip()
        else:
            # Resolve from Outline Number
            out_parts = t["outline"].split(".") if t["outline"] else []
            lote_mae = "Geral"
            lot = "Geral"
            if len(out_parts) >= 1:
                lvl1_out = out_parts[0]
                lote_mae = outline_to_name.get(lvl1_out, "Geral")
            if len(out_parts) >= 2:
                lvl2_out = ".".join(out_parts[:2])
                lot = outline_to_name.get(lvl2_out, lote_mae)
            elif len(out_parts) == 1:
                lot = lote_mae

        task_key = t["name"]
        if task_key in tasks_dict:
            task_key = f"{t['name']} (#{t['id']})"

        tasks_dict[task_key] = {
            "id": t["id"],
            "name": t["name"],
            "duration": t["duration"],
            "start": start_date.strftime("%Y-%m-%d") if start_date else None,
            "finish": finish_date.strftime("%Y-%m-%d") if finish_date else None,
            "predecessors": pred_map.get(t["id"], []),
            "successors": succ_map.get(t["id"], []),
            "lote_mae": lote_mae,
            "lot": lot,
            "is_summary": t["is_summary"],
            "outline": t.get("outline", ""),
            "wbs": t.get("wbs", ""),
            "outline_level": t.get("outline_level", 1),
            "parent_id": t.get("parent_id"),
            "daily": daily_vals
        }

    return {
        "version_name": "Importado MS Project XML",
        "tasks": tasks_dict
    }

def parse_schedule_excel_tab(excel_content: bytes) -> Dict[str, Any]:
    """
    Parses a schedule directly from an Excel file.
    Supports either:
    1. Daily distribution matrix (Piemonte 'CRONOGRAMA PROJECT' tab)
    2. Tabular task columns (ConstruKore style: ID, Nome, Início, Término, Predecessoras, Sucessoras, Lote)
    """
    wb = openpyxl.load_workbook(io.BytesIO(excel_content), read_only=True, data_only=True)
    ws_crono = None
    for s in wb.sheetnames:
        if 'cronograma' in norm(s) or 'schedule' in norm(s) or 'project' in norm(s):
            ws_crono = wb[s]
            break
    if not ws_crono:
        ws_crono = wb.active

    tasks: Dict[str, Dict[str, Any]] = {}

    # Check if this sheet is a Daily Matrix or Columnar Table
    # Sample first 10 rows
    sample_rows = list(ws_crono.iter_rows(min_row=1, max_row=10, values_only=True))
    
    # Check for daily dates row (like row 6 in Piemonte)
    has_daily_dates = False
    crono_dates = []
    daily_row_idx = None
    for r_idx, r in enumerate(sample_rows):
        d_count = sum(1 for val in r if isinstance(val, (datetime.datetime, datetime.date)))
        if d_count >= 5:
            has_daily_dates = True
            daily_row_idx = r_idx + 1
            for c_idx, val in enumerate(r):
                d = try_date(val)
                if d:
                    crono_dates.append((c_idx, d))
            break

    if has_daily_dates and crono_dates:
        # 1. Daily Matrix Mode (Standard Piemonte)
        task_counter = 1
        for row in ws_crono.iter_rows(min_row=daily_row_idx + 2, values_only=True):
            name = row[0]
            dur = row[1] if len(row) > 1 else None
            if not name:
                continue
            name_str = str(name).strip()
            daily_vals: Dict[datetime.date, float] = {}
            for c_idx, dt in crono_dates:
                if c_idx < len(row):
                    fval = try_float(row[c_idx])
                    if fval > 0:
                        daily_vals[dt] = fval
            
            d_keys = sorted([dt for dt, v in daily_vals.items() if v > 0])
            s_date = d_keys[0].strftime("%Y-%m-%d") if d_keys else None
            f_date = d_keys[-1].strftime("%Y-%m-%d") if d_keys else None

            parts = name_str.split(" - ")
            lot = " - ".join(parts[1:]).strip() if len(parts) >= 2 else "Geral"

            tasks[name_str] = {
                "id": str(task_counter),
                "name": name_str,
                "duration": try_float(dur, default=len(d_keys) or 1.0),
                "start": s_date,
                "finish": f_date,
                "predecessors": [],
                "successors": [],
                "lot": lot,
                "daily": daily_vals
            }
            task_counter += 1
    else:
        # 2. Columnar Table Mode (ConstruKore style)
        # Find header row
        header_row_idx = 1
        col_map: Dict[str, int] = {}
        for r_idx, r in enumerate(sample_rows):
            normalized = [norm(str(c or '')) for c in r]
            if any('nome' in x or 'atividade' in x or 'tarefa' in x for x in normalized):
                header_row_idx = r_idx + 1
                for c_idx, h in enumerate(normalized):
                    if any(k in h for k in ('nome', 'atividade', 'tarefa', 'package')):
                        col_map['name'] = c_idx
                    elif any(k in h for k in ('duracao', 'duration', 'dias')):
                        col_map['duration'] = c_idx
                    elif any(k in h for k in ('inicio', 'start', 'comeco')):
                        col_map['start'] = c_idx
                    elif any(k in h for k in ('termino', 'fim', 'end', 'finish')):
                        col_map['finish'] = c_idx
                    elif any(k in h for k in ('predec',)):
                        col_map['predecessors'] = c_idx
                    elif any(k in h for k in ('sucess',)):
                        col_map['successors'] = c_idx
                    elif any(k in h for k in ('lote', 'pavimento', 'local', 'grupo')):
                        col_map['lot'] = c_idx
                    elif any(k in h for k in ('id', 'codigo', 'item')):
                        col_map['id'] = c_idx
                break

        name_col = col_map.get('name', 0)
        dur_col = col_map.get('duration')
        start_col = col_map.get('start')
        finish_col = col_map.get('finish')
        pred_col = col_map.get('predecessors')
        succ_col = col_map.get('successors')
        lot_col = col_map.get('lot')
        id_col = col_map.get('id')

        for row in ws_crono.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not row or len(row) <= name_col or not row[name_col]:
                continue
            name_str = str(row[name_col]).strip()
            
            s_date = try_date(row[start_col]) if start_col is not None and start_col < len(row) else None
            f_date = try_date(row[finish_col]) if finish_col is not None and finish_col < len(row) else None
            dur = try_float(row[dur_col]) if dur_col is not None and dur_col < len(row) else 1.0

            if s_date and f_date and dur <= 1:
                dur = max(1.0, float((f_date - s_date).days + 1))

            # Build daily pct
            daily_vals = {}
            if s_date and f_date and f_date >= s_date:
                curr = s_date
                tot_d = max(1, (f_date - s_date).days + 1)
                daily_pct = 100.0 / tot_d
                while curr <= f_date:
                    daily_vals[curr] = daily_pct
                    curr += datetime.timedelta(days=1)

            preds = str(row[pred_col]).split(';') if pred_col is not None and pred_col < len(row) and row[pred_col] else []
            succs = str(row[succ_col]).split(';') if succ_col is not None and succ_col < len(row) and row[succ_col] else []
            lote_mae_col = col_map.get('lote_mae')
            lote_mae_val = str(row[lote_mae_col]).strip() if lote_mae_col is not None and lote_mae_col < len(row) and row[lote_mae_col] else None
            lot_val = str(row[lot_col]).strip() if lot_col is not None and lot_col < len(row) and row[lot_col] else None

            if not lote_mae_val or not lot_val:
                parts = name_str.split(" - ")
                if len(parts) >= 3:
                    lote_mae_val = lote_mae_val or parts[1].strip()
                    lot_val = lot_val or " - ".join(parts[2:]).strip()
                elif len(parts) == 2:
                    lote_mae_val = lote_mae_val or parts[0].strip()
                    lot_val = lot_val or parts[1].strip()
                else:
                    lote_mae_val = lote_mae_val or "Geral"
                    lot_val = lot_val or "Geral"

            t_id = str(row[id_col]).strip() if id_col is not None and id_col < len(row) and row[id_col] else str(len(tasks) + 1)
            task_key = name_str
            if task_key in tasks:
                task_key = f"{name_str} (#{t_id})"

            tasks[task_key] = {
                "id": t_id,
                "name": name_str,
                "duration": dur,
                "start": s_date.strftime("%Y-%m-%d") if s_date else None,
                "finish": f_date.strftime("%Y-%m-%d") if f_date else None,
                "predecessors": [p.strip() for p in preds if p.strip()],
                "successors": [s.strip() for s in succs if s.strip()],
                "lote_mae": lote_mae_val,
                "lot": lot_val,
                "daily": daily_vals
            }

    wb.close()
    return {
        "version_name": "Importado Planilha Excel",
        "tasks": tasks
    }

def parse_ms_project_mpp(mpp_content: bytes, filename: str = "cronograma.mpp") -> Dict[str, Any]:
    """
    Parses a native MS Project (.mpp) binary file using MPXJ and portable OpenJDK 17.
    Extracts all tasks, durations, start/finish dates, outline codes, predecessors, lote_mae and lot.
    """
    _ensure_jvm()
    UniversalProjectReader = jpype.JClass("org.mpxj.reader.UniversalProjectReader")
    reader = UniversalProjectReader()

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".mpp", delete=False) as tmp:
            tmp.write(mpp_content)
            tmp_path = tmp.name

        project = reader.read(tmp_path)
        mpxj_tasks = project.getTasks()

        tasks_dict: Dict[str, Dict[str, Any]] = {}
        pred_map: Dict[str, List[str]] = {}
        succ_map: Dict[str, List[str]] = {}
        raw_tasks = []

        type_map = {
            'FINISH_START': 'FS',
            'START_START': 'SS',
            'FINISH_FINISH': 'FF',
            'START_FINISH': 'SF'
        }

        def get_mpxj_hierarchy(task):
            parents = []
            curr = task.getParentTask()
            while curr:
                if curr.getID() != 0:
                    p_name = str(curr.getName() or '').strip()
                    if p_name:
                        parents.append(p_name)
                curr = curr.getParentTask()
            parents.reverse()
            return parents

        for i in range(mpxj_tasks.size()):
            t = mpxj_tasks.get(i)
            if not t or not t.getName():
                continue
            t_id = str(t.getID())
            if t_id == "0":
                continue
            name = str(t.getName()).strip()
            outline = str(t.getOutlineNumber()) if t.getOutlineNumber() else ""
            wbs = str(t.getWBS()) if t.getWBS() else outline
            try:
                outline_level = int(t.getOutlineLevel()) if t.getOutlineLevel() is not None else (len(wbs.split('.')) if wbs else 1)
            except Exception:
                outline_level = len(wbs.split('.')) if wbs else 1
            parent = t.getParentTask()
            parent_id = str(parent.getID()) if parent and parent.getID() != 0 else None
            is_summary = bool(t.getSummary())
            
            dur_obj = t.getDuration()
            dur_days = float(dur_obj.getDuration()) if dur_obj else 0.0
            
            s_date = try_date(t.getStart())
            f_date = try_date(t.getFinish())

            preds = []
            if t.getPredecessors():
                for r in t.getPredecessors():
                    pred_t = r.getPredecessorTask()
                    if pred_t:
                        rtype = type_map.get(str(r.getType()), 'FS')
                        lag_val = 0
                        if r.getLag():
                            try:
                                lag_val = int(r.getLag().getDuration())
                            except Exception:
                                pass
                        lag_str = f'+{lag_val}d' if lag_val > 0 else (f'{lag_val}d' if lag_val < 0 else '')
                        preds.append(f'{pred_t.getID()}{rtype}{lag_str}')

            pred_map[t_id] = preds

            # Derive lote_mae and lot
            parts = name.split(" - ")
            if len(parts) >= 3:
                lote_mae = parts[1].strip()
                lot = " - ".join(parts[2:]).strip()
            elif len(parts) == 2:
                lote_mae = parts[0].strip()
                lot = parts[1].strip()
            else:
                parents = get_mpxj_hierarchy(t)
                lote_mae = parents[0] if len(parents) >= 1 else "Geral"
                lot = parents[1] if len(parents) >= 2 else (parents[0] if parents else "Geral")

            raw_tasks.append({
                "id": t_id,
                "name": name,
                "duration": dur_days,
                "start": s_date,
                "finish": f_date,
                "outline": wbs,
                "wbs": wbs,
                "outline_level": outline_level,
                "parent_id": parent_id,
                "lote_mae": lote_mae,
                "lot": lot,
                "is_summary": is_summary
            })

        # Build successor map
        for t in raw_tasks:
            succ_map[t["id"]] = []

        for t in raw_tasks:
            for p_str in pred_map.get(t["id"], []):
                pid_num = ''.join(c for c in p_str.split('F')[0].split('S')[0] if c.isdigit())
                if pid_num in succ_map:
                    succ_map[pid_num].append(t["id"])

        for t in raw_tasks:
            s_date = t["start"]
            f_date = t["finish"]
            daily_vals = {}
            if s_date and f_date and f_date >= s_date:
                curr = s_date
                tot_d = max(1, (f_date - s_date).days + 1)
                daily_pct = 100.0 / tot_d
                while curr <= f_date:
                    if curr.weekday() < 5:
                        daily_vals[curr] = daily_pct
                    else:
                        daily_vals[curr] = 0.0
                    curr += datetime.timedelta(days=1)
                
                active_days = sum(1 for v in daily_vals.values() if v > 0)
                if active_days > 0:
                    pct_per_day = 100.0 / active_days
                    for d in daily_vals:
                        if daily_vals[d] > 0:
                            daily_vals[d] = pct_per_day

            task_key = t["name"]
            if task_key in tasks_dict:
                task_key = f"{t['name']} (#{t['id']})"

            tasks_dict[task_key] = {
                "id": t["id"],
                "name": t["name"],
                "duration": t["duration"],
                "start": s_date.strftime("%Y-%m-%d") if s_date else None,
                "finish": f_date.strftime("%Y-%m-%d") if f_date else None,
                "predecessors": pred_map.get(t["id"], []),
                "successors": succ_map.get(t["id"], []),
                "lote_mae": t["lote_mae"],
                "lot": t["lot"],
                "is_summary": t["is_summary"],
                "outline": t.get("outline", ""),
                "wbs": t.get("wbs", ""),
                "outline_level": t.get("outline_level", 1),
                "parent_id": t.get("parent_id"),
                "daily": daily_vals
            }

        return {
            "version_name": f"Importado MS Project (.mpp)",
            "tasks": tasks_dict
        }
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass


def get_group_leaf_tasks(group_id_or_name: str, tasks: Dict[str, Any]) -> List[Dict[str, Any]]:
    """
    Given a group task ID or name:
    If it's an individual task (is_summary == False), returns [task].
    If it's a summary task (is_summary == True), finds all descendant leaf tasks (is_summary == False).
    """
    target = None
    target_id = str(group_id_or_name).strip()

    # Match by ID, name or key
    for k, t in tasks.items():
        if str(t.get("id")) == target_id or t.get("name") == group_id_or_name or k == group_id_or_name:
            target = t
            break

    if not target:
        return []

    if not target.get("is_summary", False):
        return [target]

    target_outline = str(target.get("outline") or target.get("wbs") or "")
    target_id = str(target.get("id") or "")

    leaves = []
    # 1. Match by WBS/Outline prefix
    if target_outline:
        prefix = target_outline + "."
        for k, t in tasks.items():
            t_out = str(t.get("outline") or t.get("wbs") or "")
            if t_out.startswith(prefix) and not t.get("is_summary", False):
                leaves.append(t)

    # 2. Fallback by parent_id chain or name prefix if WBS didn't match
    if not leaves:
        def is_descendant(t_check, parent_id_target):
            curr_pid = t_check.get("parent_id")
            visited = set()
            while curr_pid and curr_pid not in visited:
                if str(curr_pid) == str(parent_id_target):
                    return True
                visited.add(curr_pid)
                p_task = next((x for x in tasks.values() if str(x.get("id")) == str(curr_pid)), None)
                curr_pid = p_task.get("parent_id") if p_task else None
            return False

        for k, t in tasks.items():
            if not t.get("is_summary", False):
                if is_descendant(t, target_id):
                    leaves.append(t)
                elif target.get("name") and (target["name"].lower() in (t.get("lote_mae", "") or "").lower() or target["name"].lower() in t.get("name", "").lower()):
                    leaves.append(t)

    # Ensure no duplicates by ID
    unique_leaves = []
    seen_ids = set()
    for leaf in leaves:
        lid = leaf.get("id") or leaf.get("name")
        if lid not in seen_ids:
            seen_ids.add(lid)
            unique_leaves.append(leaf)

    return unique_leaves

