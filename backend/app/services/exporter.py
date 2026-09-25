import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
from typing import Dict, Any, Tuple, Optional

# --- Color Palettes & Styles ---
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_fill_accent = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
past_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
future_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
total_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

lvl1_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
lvl2_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
lvl3_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
lvl6_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
bold_font = Font(name="Segoe UI", size=10, bold=True)
normal_font = Font(name="Segoe UI", size=10)
italic_font = Font(name="Segoe UI", size=9, italic=True, color="475569")
title_font = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

def _auto_fit_columns(ws, min_widths=None):
    ws.views.sheetView[0].showGridLines = True
    min_widths = min_widths or {}
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = min_widths.get(col_letter, 10)
        for cell in col[:120]:
            val = cell.value
            if val is not None:
                val_str = str(val)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)


def add_sheet_summary(wb: openpyxl.Workbook, replan_result: Dict[str, Any]):
    ws = wb.create_sheet(title="Resumo Curva S")
    ws["A1"] = f"REPLANEJAMENTO FÍSICO-FINANCEIRO - {replan_result.get('project_name', 'OBRA').upper()}"
    ws["A1"].font = title_font

    ws["A2"] = (
        f"Data de Corte: {replan_result.get('cutoff_date', '-')} (Medição {replan_result.get('cutoff_med_num', '-')}) | "
        f"Ciclo: Dia {replan_result.get('cycle_config', {}).get('start_day', 21)} ao {replan_result.get('cycle_config', {}).get('end_day', 20)} | "
        f"Orçamento Total: R$ {replan_result.get('total_budget', 0):,.2f} | "
        f"Realizado: {replan_result.get('total_measured_pct', 0)}% | "
        f"Saldo: {replan_result.get('total_balance_pct', 0)}%"
    )
    ws["A2"].font = italic_font

    headers = [
        "Mês", "Período", "Início Ciclo", "Fim Ciclo", "Tipo",
        "R$ Mês Replanejado", "% Mês", "R$ Acumulado", "% Acumulado",
        "R$ Mês Realizado", "% Acumulado Realizado"
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cycles = replan_result.get("cycles", [])
    s_curve = replan_result.get("s_curve", {})

    for row_idx, c in enumerate(cycles, 5):
        idx = row_idx - 5
        tipo = "REALIZADO" if c.get("is_past") else "REPLANEJADO"
        fill = past_fill if c.get("is_past") else future_fill

        ws.cell(row=row_idx, column=1, value=c.get("name")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=c.get("period_label")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3, value=c.get("start")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=c.get("end")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=5, value=tipo).alignment = Alignment(horizontal="center")

        rep_m_val = s_curve.get("replan_month_vals", [])[idx] if idx < len(s_curve.get("replan_month_vals", [])) else 0.0
        rep_m_pct = s_curve.get("replan_month_pcts", [])
        m_pct_val = (rep_m_pct[idx] / 100.0) if idx < len(rep_m_pct) and rep_m_pct[idx] is not None else 0.0

        c6 = ws.cell(row=row_idx, column=6, value=rep_m_val or 0.0)
        c6.number_format = 'R$ #,##0.00'
        c7 = ws.cell(row=row_idx, column=7, value=m_pct_val or 0.0)
        c7.number_format = '0.00%'

        rep_ac_val = s_curve.get("replan_accum_vals", [])
        ac_val = rep_ac_val[idx] if idx < len(rep_ac_val) and rep_ac_val[idx] is not None else 0.0
        c8 = ws.cell(row=row_idx, column=8, value=ac_val)
        c8.number_format = 'R$ #,##0.00'

        rep_ac_pct = s_curve.get("replan_accum_pcts", [])
        ac_pct = (rep_ac_pct[idx] / 100.0) if idx < len(rep_ac_pct) and rep_ac_pct[idx] is not None else 0.0
        c9 = ws.cell(row=row_idx, column=9, value=ac_pct)
        c9.number_format = '0.00%'

        real_m = s_curve.get("realized_month_vals", [])[idx] if idx < len(s_curve.get("realized_month_vals", [])) else None
        real_acum = s_curve.get("realized_accum_pcts", [])[idx] if idx < len(s_curve.get("realized_accum_pcts", [])) else None

        c10 = ws.cell(row=row_idx, column=10, value=real_m if real_m is not None else "-")
        if real_m is not None:
            c10.number_format = 'R$ #,##0.00'
            c10.alignment = Alignment(horizontal="right")
        else:
            c10.alignment = Alignment(horizontal="center")

        c11 = ws.cell(row=row_idx, column=11, value=(real_acum / 100.0) if real_acum is not None else "-")
        if real_acum is not None:
            c11.number_format = '0.00%'
            c11.alignment = Alignment(horizontal="right")
        else:
            c11.alignment = Alignment(horizontal="center")

        for col_idx in range(1, 12):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = normal_font
            cell.border = thin_border
            if col_idx in (5, 6, 7, 8, 9):
                cell.fill = fill

    _auto_fit_columns(ws)


def add_sheet_ff_replan(wb: openpyxl.Workbook, replan_result: Dict[str, Any], view_mode: str = "val"):
    title = "FF Replanejado (%)" if view_mode == "pct" else "FF Replanejado (R$)"
    ws = wb.create_sheet(title=title)
    cycles = replan_result.get("cycles", [])
    total_budget = replan_result.get("total_budget", 1.0) or 1.0

    base_headers = ["Nível", "Código", "Descrição", "Und", "Qtd", "Orçamento Total (R$)", "% Medido", "R$ Medido", "Saldo a Replan. (R$)"]
    suffix = " (%)" if view_mode == "pct" else " (R$)"
    cycle_headers = [c["period_label"] + suffix for c in cycles]
    total_col_header = "Total Replanejado (%)" if view_mode == "pct" else "Total Replanejado (R$)"
    total_headers = base_headers + cycle_headers + [total_col_header]

    ws.append(total_headers)
    for col_idx in range(1, len(total_headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    tree_rows = replan_result.get("tab_ff_replanejado") or replan_result.get("tree_rows", [])
    for row_idx, node in enumerate(tree_rows, 2):
        lvl = node["level"]
        is_bold = (lvl <= 3)
        indent = "  " * (lvl - 1)

        r_cells = [
            lvl,
            node["code"],
            indent + str(node["description"] or ""),
            node.get("unit") or "",
            node.get("quantity") or 0.0,
            node.get("budget", 0.0),
            (node.get("med_pct", 0.0) / 100.0),
            node.get("med_val", 0.0),
            node.get("saldo", 0.0),
        ]

        row_total = 0.0
        for c in cycles:
            val_r = node["cycle_vals"].get(str(c["num"]), 0.0)
            row_total += val_r
            if view_mode == "pct":
                # Percentage relative to total project budget
                pct_val = (val_r / total_budget)
                r_cells.append(pct_val)
            else:
                r_cells.append(val_r)

        if view_mode == "pct":
            r_cells.append(row_total / total_budget)
        else:
            r_cells.append(row_total)

        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[1].alignment = Alignment(horizontal="left")
        cur_row[3].alignment = Alignment(horizontal="center")
        cur_row[4].number_format = '#,##0.00'
        cur_row[5].number_format = 'R$ #,##0.00'
        cur_row[6].number_format = '0.00%'
        cur_row[7].number_format = 'R$ #,##0.00'
        cur_row[8].number_format = 'R$ #,##0.00'

        for c_offset, c in enumerate(cycles):
            c_cell = cur_row[9 + c_offset]
            if view_mode == "pct":
                c_cell.number_format = '0.00%'
            else:
                c_cell.number_format = 'R$ #,##0.00'
            if c.get("is_past"):
                c_cell.fill = past_fill

        cur_row[-1].number_format = '0.00%' if view_mode == "pct" else 'R$ #,##0.00'
        cur_row[-1].fill = total_fill

        if is_bold:
            for cell in cur_row:
                cell.font = bold_font
                if lvl == 1:
                    cell.fill = lvl1_fill
                elif lvl == 2:
                    cell.fill = lvl2_fill
                elif lvl == 3:
                    cell.fill = lvl3_fill

    _auto_fit_columns(ws, min_widths={'C': 38})


def add_sheet_distribuicao(wb: openpyxl.Workbook, replan_result: Dict[str, Any], view_mode: str = "val"):
    if view_mode == "pct":
        mode_label = "% Obra"
    elif view_mode == "pct_service":
        mode_label = "Peso Serviço"
    else:
        mode_label = "R$"

    ws = wb.create_sheet(title=f"Distribuição ({mode_label})")
    cycles = replan_result.get("cycles", [])
    total_budget = replan_result.get("total_budget", 1.0) or 1.0

    alloc_header = "Alocado (% Obra)" if view_mode == "pct" else ("Peso (%)" if view_mode == "pct_service" else "Alocado (R$)")
    tot_header = "Total (% Obra)" if view_mode == "pct" else ("Total (%)" if view_mode == "pct_service" else "Total (R$)")
    base_headers = ["Nível", "Código", "Descrição / Atividade", "Duração (d)", alloc_header]
    cycle_headers = [f"{c['period_label']} ({mode_label})" for c in cycles]
    total_headers = base_headers + cycle_headers + [tot_header]

    ws.append(total_headers)
    for col_idx in range(1, len(total_headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill_accent
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = replan_result.get("tab_distribuicao", [])
    for row_idx, r in enumerate(rows, 2):
        lvl = r.get("level", 5)
        indent = "    " if lvl == 6 else ""
        desc = indent + str(r.get("description") or "")
        dur = r.get("duration", 0.0)

        # Allocated value / pct
        alloc_val_r = r.get("allocated_val", 0.0)
        if view_mode == "pct":
            alloc_cell_val = (alloc_val_r / total_budget)
        elif view_mode == "pct_service":
            alloc_cell_val = (r.get("weight_pct", 0.0) / 100.0)
        else:
            alloc_cell_val = alloc_val_r

        r_cells = [
            lvl,
            r.get("code", ""),
            desc,
            dur if dur > 0 else "-",
            alloc_cell_val
        ]

        row_tot = 0.0
        for c in cycles:
            c_str = str(c["num"])
            v_reais = r.get("cycle_vals", {}).get(c_str, 0.0)
            row_tot += v_reais

            if view_mode == "pct":
                c_val = (v_reais / total_budget)
            elif view_mode == "pct_service":
                c_val = (v_reais / alloc_val_r) if alloc_val_r > 0 else 0.0
            else:
                c_val = v_reais
            r_cells.append(c_val)

        if view_mode == "pct":
            r_cells.append(row_tot / total_budget)
        elif view_mode == "pct_service":
            r_cells.append((row_tot / alloc_val_r) if alloc_val_r > 0 else 0.0)
        else:
            r_cells.append(row_tot)

        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[1].alignment = Alignment(horizontal="left")
        cur_row[3].alignment = Alignment(horizontal="center")

        # Format alloc cell
        if view_mode in ("pct", "pct_service"):
            cur_row[4].number_format = '0.00%'
        else:
            cur_row[4].number_format = 'R$ #,##0.00'

        # Format cycle cells
        for c_offset, c in enumerate(cycles):
            c_cell = cur_row[5 + c_offset]
            if view_mode in ("pct", "pct_service"):
                c_cell.number_format = '0.00%'
            else:
                c_cell.number_format = 'R$ #,##0.00'
            if c.get("is_past"):
                c_cell.fill = past_fill

        cur_row[-1].number_format = '0.00%' if view_mode in ("pct", "pct_service") else 'R$ #,##0.00'
        cur_row[-1].fill = total_fill

        if lvl == 5:
            for cell in cur_row:
                cell.font = bold_font
                cell.fill = lvl2_fill
        elif lvl == 6:
            for cell in cur_row:
                cell.font = normal_font
                if not cell.fill or cell.fill.fill_type is None:
                    cell.fill = lvl6_fill

    _auto_fit_columns(ws, min_widths={'C': 42})


def add_sheet_ff_baseline(wb: openpyxl.Workbook, replan_result: Dict[str, Any], view_mode: str = "val"):
    title = "FF Previsto Base (%)" if view_mode == "pct" else "FF Previsto Base (R$)"
    ws = wb.create_sheet(title=title)
    cycles = replan_result.get("cycles", [])
    total_budget = replan_result.get("total_budget", 1.0) or 1.0

    base_headers = ["Nível", "Código", "Descrição", "Und", "Qtd", "Orçamento Previsto (R$)"]
    suffix = " (%)" if view_mode == "pct" else " (R$)"
    cycle_headers = [c["period_label"] + suffix for c in cycles]
    total_col_header = "Total Previsto (%)" if view_mode == "pct" else "Total Previsto (R$)"
    total_headers = base_headers + cycle_headers + [total_col_header]

    ws.append(total_headers)
    for col_idx in range(1, len(total_headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = replan_result.get("tab_fisico_financeiro", [])
    for row_idx, node in enumerate(rows, 2):
        lvl = node["level"]
        is_bold = (lvl <= 3)
        indent = "  " * (lvl - 1)

        r_cells = [
            lvl,
            node["code"],
            indent + str(node["description"] or ""),
            node.get("unit") or "",
            node.get("quantity") or 0.0,
            node.get("total_price", node.get("budget", 0.0)),
        ]

        row_tot = 0.0
        for c in cycles:
            val_r = node.get("cycle_vals", {}).get(str(c["num"]), 0.0)
            row_tot += val_r
            if view_mode == "pct":
                r_cells.append(val_r / total_budget)
            else:
                r_cells.append(val_r)

        if view_mode == "pct":
            r_cells.append(row_tot / total_budget)
        else:
            r_cells.append(row_tot)

        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[1].alignment = Alignment(horizontal="left")
        cur_row[3].alignment = Alignment(horizontal="center")
        cur_row[4].number_format = '#,##0.00'
        cur_row[5].number_format = 'R$ #,##0.00'

        for c_offset, c in enumerate(cycles):
            c_cell = cur_row[6 + c_offset]
            c_cell.number_format = '0.00%' if view_mode == "pct" else 'R$ #,##0.00'

        cur_row[-1].number_format = '0.00%' if view_mode == "pct" else 'R$ #,##0.00'
        cur_row[-1].fill = total_fill

        if is_bold:
            for cell in cur_row:
                cell.font = bold_font
                if lvl == 1:
                    cell.fill = lvl1_fill
                elif lvl == 2:
                    cell.fill = lvl2_fill

    _auto_fit_columns(ws, min_widths={'C': 38})


def add_sheet_replanejado(wb: openpyxl.Workbook, replan_result: Dict[str, Any], view_mode: str = "val"):
    title = "Replanejado Saldo (%)" if view_mode == "pct" else "Replanejado Saldo (R$)"
    ws = wb.create_sheet(title=title)
    cycles = replan_result.get("cycles", [])
    total_budget = replan_result.get("total_budget", 1.0) or 1.0

    base_headers = ["Nível", "Código", "Descrição / Atividade", "Peso Futuro (%)", "Saldo Alocado (R$)"]
    suffix = " (%)" if view_mode == "pct" else " (R$)"
    cycle_headers = [c["period_label"] + suffix for c in cycles]
    total_headers = base_headers + cycle_headers + ["Total Saldo Distribuído" + suffix]

    ws.append(total_headers)
    for col_idx in range(1, len(total_headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill_accent
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = replan_result.get("tab_replanejado", [])
    for row_idx, r in enumerate(rows, 2):
        lvl = r.get("level", 5)
        indent = "    " if lvl == 6 else ""
        desc = indent + str(r.get("description") or "")
        w_fut = (r.get("future_weight_pct", 0.0) / 100.0)
        saldo_aloc = r.get("allocated_saldo", 0.0)

        r_cells = [
            lvl,
            r.get("code", ""),
            desc,
            w_fut,
            saldo_aloc
        ]

        row_tot = 0.0
        for c in cycles:
            val_r = r.get("cycle_vals", {}).get(str(c["num"]), 0.0)
            row_tot += val_r
            if view_mode == "pct":
                r_cells.append(val_r / total_budget)
            else:
                r_cells.append(val_r)

        if view_mode == "pct":
            r_cells.append(row_tot / total_budget)
        else:
            r_cells.append(row_tot)

        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[1].alignment = Alignment(horizontal="left")
        cur_row[3].number_format = '0.00%'
        cur_row[4].number_format = 'R$ #,##0.00'

        for c_offset, c in enumerate(cycles):
            c_cell = cur_row[5 + c_offset]
            c_cell.number_format = '0.00%' if view_mode == "pct" else 'R$ #,##0.00'
            if c.get("is_past"):
                c_cell.fill = past_fill

        cur_row[-1].number_format = '0.00%' if view_mode == "pct" else 'R$ #,##0.00'
        cur_row[-1].fill = total_fill

        if lvl == 5:
            for cell in cur_row:
                cell.font = bold_font
                cell.fill = lvl2_fill
        elif lvl == 6:
            for cell in cur_row:
                cell.font = normal_font

    _auto_fit_columns(ws, min_widths={'C': 40})


def add_sheet_medicao(wb: openpyxl.Workbook, replan_result: Dict[str, Any]):
    ws = wb.create_sheet(title="Medição da Obra")
    cycles = replan_result.get("cycles", [])

    base_headers = ["Nível", "Código", "Descrição", "Orçamento Total (R$)", "% Medido Acumulado", "R$ Medido Acumulado", "Saldo Restante (R$)"]
    cycle_headers = [f"{c['name']} ({c['period_label']}) %" for c in cycles]
    total_headers = base_headers + cycle_headers

    ws.append(total_headers)
    for col_idx in range(1, len(total_headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = replan_result.get("tab_medicao", [])
    for row_idx, r in enumerate(rows, 2):
        lvl = r.get("level", 5)
        indent = "  " * (lvl - 1)
        b_val = r.get("total_price", 0.0)
        acc_pct = (r.get("accum_realizado_pct", 0.0) / 100.0)
        acc_val = r.get("accum_realizado_val", 0.0)
        saldo = r.get("saldo", 0.0)

        r_cells = [
            lvl,
            r.get("code", ""),
            indent + str(r.get("description") or ""),
            b_val,
            acc_pct,
            acc_val,
            saldo
        ]

        m_map = r.get("monthly_map", {})
        for c in cycles:
            val_pct = (m_map.get(str(c["num"]), 0.0) / 100.0)
            r_cells.append(val_pct)

        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[1].alignment = Alignment(horizontal="left")
        cur_row[3].number_format = 'R$ #,##0.00'
        cur_row[4].number_format = '0.00%'
        cur_row[5].number_format = 'R$ #,##0.00'
        cur_row[6].number_format = 'R$ #,##0.00'

        for c_offset, c in enumerate(cycles):
            c_cell = cur_row[7 + c_offset]
            c_cell.number_format = '0.00%'
            if c.get("is_past"):
                c_cell.fill = past_fill

        if lvl <= 3:
            for cell in cur_row:
                cell.font = bold_font
                if lvl == 1:
                    cell.fill = lvl1_fill
                elif lvl == 2:
                    cell.fill = lvl2_fill

    _auto_fit_columns(ws, min_widths={'C': 38})


def add_sheet_orcamento(wb: openpyxl.Workbook, replan_result: Dict[str, Any]):
    ws = wb.create_sheet(title="Orçamento EAP")
    headers = ["Nível", "Código", "Descrição", "Und", "Qtd", "Preço Unitário (R$)", "Preço Total (R$)", "% Participação"]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = replan_result.get("tab_orcamento", [])
    for row_idx, r in enumerate(rows, 2):
        lvl = r.get("level", 5)
        indent = "  " * (lvl - 1)
        r_cells = [
            lvl,
            r.get("code", ""),
            indent + str(r.get("description") or ""),
            r.get("unit") or "",
            r.get("quantity", 0.0),
            r.get("unit_price", 0.0),
            r.get("total_price", 0.0),
            (r.get("pct_share", 0.0) / 100.0)
        ]
        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[1].alignment = Alignment(horizontal="left")
        cur_row[3].alignment = Alignment(horizontal="center")
        cur_row[4].number_format = '#,##0.00'
        cur_row[5].number_format = 'R$ #,##0.00'
        cur_row[6].number_format = 'R$ #,##0.00'
        cur_row[7].number_format = '0.00%'

        if lvl <= 3:
            for cell in cur_row:
                cell.font = bold_font
                if lvl == 1:
                    cell.fill = lvl1_fill
                elif lvl == 2:
                    cell.fill = lvl2_fill

    _auto_fit_columns(ws, min_widths={'C': 40})


def add_sheet_cronograma(wb: openpyxl.Workbook, replan_result: Dict[str, Any]):
    ws = wb.create_sheet(title="Cronograma de Atividades")
    headers = ["ID", "Nome da Atividade", "Duração (dias)", "Início", "Término", "Predecessoras", "Lote Mãe", "Lote / Pavimento", "Situação"]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill_accent
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    tasks = replan_result.get("tab_cronograma", {}).get("tasks", [])
    for row_idx, t in enumerate(tasks, 2):
        preds = ", ".join(t.get("predecessors", [])) if isinstance(t.get("predecessors"), list) else str(t.get("predecessors", ""))
        sit = "Futura (Replanejada)" if t.get("is_future") else "Realizada / Passada"

        r_cells = [
            t.get("id", ""),
            t.get("name", ""),
            t.get("duration", 0),
            t.get("start_date", "-"),
            t.get("end_date", "-"),
            preds or "-",
            t.get("lote_mae", "Geral"),
            t.get("lot", "Geral"),
            sit
        ]
        ws.append(r_cells)
        cur_row = ws[row_idx]
        cur_row[0].alignment = Alignment(horizontal="center")
        cur_row[2].alignment = Alignment(horizontal="center")
        cur_row[3].alignment = Alignment(horizontal="center")
        cur_row[4].alignment = Alignment(horizontal="center")
        cur_row[8].alignment = Alignment(horizontal="center")
        if t.get("is_future"):
            cur_row[8].fill = future_fill
        else:
            cur_row[8].fill = past_fill

    _auto_fit_columns(ws, min_widths={'B': 35})


def export_replan_to_excel(
    replan_result: Dict[str, Any],
    export_target: str = "current",
    active_tab: str = "ff_replanejado",
    dist_view_mode: str = "val",
    ff_rep_view_mode: str = "val",
    ff_view_mode: str = "val"
) -> Tuple[io.BytesIO, str]:
    """
    Exports a high-fidelity Excel workbook based on user parameters.
    - If export_target == 'current': exports the exact tab currently open with its view mode (% or R$)
    - If export_target == 'all': exports all 8 application tabs in a complete consolidated workbook
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    proj_name = replan_result.get("project_name", "PLATEA").replace(" ", "_").upper()
    cutoff_med = replan_result.get("cutoff_med_num", 11)

    if export_target == "all":
        # Consolidated workbook with all modules
        add_sheet_summary(wb, replan_result)
        add_sheet_orcamento(wb, replan_result)
        add_sheet_distribuicao(wb, replan_result, view_mode=dist_view_mode)
        add_sheet_ff_baseline(wb, replan_result, view_mode=ff_view_mode)
        add_sheet_medicao(wb, replan_result)
        add_sheet_replanejado(wb, replan_result, view_mode=ff_rep_view_mode)
        add_sheet_ff_replan(wb, replan_result, view_mode=ff_rep_view_mode)
        add_sheet_cronograma(wb, replan_result)
        filename = f"{proj_name}_Pasta_Completa_M{cutoff_med}.xlsx"
    else:
        # Export open tab with its active display mode
        add_sheet_summary(wb, replan_result)  # Executive summary is always included for context

        if active_tab == "distribuicao":
            add_sheet_distribuicao(wb, replan_result, view_mode=dist_view_mode)
            mode_tag = "Percentual_Obra" if dist_view_mode == "pct" else ("Peso_Servico" if dist_view_mode == "pct_service" else "Valor_R$")
            filename = f"{proj_name}_Distribuicao_{mode_tag}_M{cutoff_med}.xlsx"
        elif active_tab == "ff_replanejado":
            add_sheet_ff_replan(wb, replan_result, view_mode=ff_rep_view_mode)
            mode_tag = "Percentual" if ff_rep_view_mode == "pct" else "Valor_R$"
            filename = f"{proj_name}_FF_Replanejado_{mode_tag}_M{cutoff_med}.xlsx"
        elif active_tab == "fisico_financeiro":
            add_sheet_ff_baseline(wb, replan_result, view_mode=ff_view_mode)
            mode_tag = "Percentual" if ff_view_mode == "pct" else "Valor_R$"
            filename = f"{proj_name}_FF_Previsto_{mode_tag}_M{cutoff_med}.xlsx"
        elif active_tab == "replanejado":
            add_sheet_replanejado(wb, replan_result, view_mode=ff_rep_view_mode)
            mode_tag = "Percentual" if ff_rep_view_mode == "pct" else "Valor_R$"
            filename = f"{proj_name}_Replanejado_Saldo_{mode_tag}_M{cutoff_med}.xlsx"
        elif active_tab == "medicao":
            add_sheet_medicao(wb, replan_result)
            filename = f"{proj_name}_Medicao_Obra_M{cutoff_med}.xlsx"
        elif active_tab == "orcamento":
            add_sheet_orcamento(wb, replan_result)
            filename = f"{proj_name}_Orcamento_EAP_M{cutoff_med}.xlsx"
        elif active_tab == "cronograma":
            add_sheet_cronograma(wb, replan_result)
            filename = f"{proj_name}_Cronograma_Atividades_M{cutoff_med}.xlsx"
        else:  # dashboard
            add_sheet_ff_replan(wb, replan_result, view_mode=ff_rep_view_mode)
            filename = f"{proj_name}_Resumo_Replanejamento_M{cutoff_med}.xlsx"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output, filename
